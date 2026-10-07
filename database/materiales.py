"""Solicitudes persistentes con autorización, historial y avisos en una transacción."""
import json
import math
from datetime import datetime
from decimal import Decimal

from sqlalchemy import select, or_
from sqlalchemy.orm import Session
from sqlalchemy.orm.exc import StaleDataError
from database.conexion import engine
from database.modelos import NotificacionInterna, UsuarioRootMine
from database.modelos_materiales import SolicitudMaterial
from database.usuarios import _a_dict, _areas_responsabilidad, _norm, _resolver_responsable

ANALISTA_FIJO = "rfernandezc@agrosuper.com"
SUBGERENTE_FIJO = "cvelasquez@agrosuper.com"

PERFILES = ("Contrapedido", "Stock de seguridad", "Pronóstico")
PENDIENTE_MRP = "Pendiente exportación MRP"
LISTO_MRP = "Listo · exportado MRP"
LISTO_SS = "Listo · stock de seguridad"
ESTADOS = ("Borrador", "Pendiente jefe", "Pendiente análisis", "Pendiente subgerente", "Pendiente carga SAP", "Cargado en SAP", "Devuelta", "Rechazada", PENDIENTE_MRP, LISTO_MRP, LISTO_SS)


def criticidad(m_rotativo, smac, co, te, ma, estrategico):
    if m_rotativo:
        return "M"
    if any(x not in (1, 2, 3) for x in (smac, co, te, ma)):
        raise ValueError("Selecciona todas las respuestas de criticidad.")
    # Árbol SV-DC-MAN-019, revisión de mayo 2025.
    if smac == 1 or co == 1 or (te == 1 and ma == 1):
        return "Z" if estrategico else "A"
    if te == 3 or ma == 3:
        return "C"
    return "B"


def impacto(d):
    def n(k):
        v = Decimal(str(d.get(k, 0)))
        if not v.is_finite() or v < 0:
            raise ValueError(f"{k}: debe ser un número no negativo.")
        return v
    ss = n("ss_propuesto")
    precio = n("precio")
    stock = n("stock")
    consumo = n("consumo_mensual")
    incremento = max(Decimal(0), ss - n("ss_actual")) * precio
    compra = max(Decimal(0), ss + n("reservas") - stock - n("oc")) * precio
    bodega = n("valor_bodega")
    total_ss = ss * precio
    proyectada = bodega + compra
    cobertura_ss = None
    cobertura_stock = stock / consumo if consumo else None
    promedio = n("stock_promedio") / consumo if consumo and d.get("stock_promedio_observado") else None
    uso = n("uso_ss_cantidad") if d.get("uso_ss_registrado") else None
    ss_referencia = n("ss_historico_referencia") if d.get("uso_ss_registrado") else Decimal(0)
    periodo = n("periodo_consumo_meses")
    consumo_periodo = consumo * periodo
    consumo_emergencia = uso / periodo if uso is not None and periodo else None
    cobertura_ss = ss / consumo_emergencia if consumo_emergencia else None
    return {"incremento_ss": float(incremento), "reposicion": float(compra),
            "porcentaje": float(incremento / bodega * 100) if bodega else None,
            "bodega_proyectada": float(proyectada),
            "porcentaje_reposicion": float(compra / bodega * 100) if bodega else None,
            "bodega_con_material_solicitado": float(bodega + total_ss),
            "valor_ss_total": float(total_ss), "valor_stock_material_actual": float(stock * precio),
            "porcentaje_ss_total_bodega_actual": float(total_ss / bodega * 100) if bodega else None,
            "porcentaje_ss_total_bodega_proyectada": float(total_ss / proyectada * 100) if proyectada else None,
            "cobertura_ss_meses": float(cobertura_ss) if cobertura_ss is not None else None,
            "cobertura_ss_dias": float(cobertura_ss * 30) if cobertura_ss is not None else None,
            "cobertura_stock_meses": float(cobertura_stock) if cobertura_stock is not None else None,
            "permanencia_promedio_meses": float(promedio) if promedio is not None else None,
            "permanencia_promedio_dias": float(promedio * 30) if promedio is not None else None,
            "uso_ss_cantidad": float(uso) if uso is not None else None,
            "consumo_emergencia_ss_mensual": float(consumo_emergencia) if consumo_emergencia is not None else None,
            "uso_ss_valorizado": float(uso * precio) if uso is not None else None,
            "uso_ss_porcentaje_consumo_periodo": float(uso / consumo_periodo * 100) if uso is not None and consumo_periodo else None,
            "uso_ss_reservas_equivalentes": float(uso / ss_referencia) if uso is not None and ss_referencia else None}


def _usuario(s, correo):
    reg = s.scalar(select(UsuarioRootMine).where(UsuarioRootMine.correo == correo.strip().lower(), UsuarioRootMine.activo.is_(True)))
    if not reg:
        raise ValueError("La cuenta no está activa.")
    return _a_dict(reg)


def candidatos(usuarios, etapa, centro, area):
    return [u for u in usuarios if u.get("activo", True) and elegible(u, etapa, centro, area)]


def elegible(u, etapa, centro, area):
    mismo = str(u.get("centro", "")) == str(centro)
    if etapa == "analista":
        return (u.get("es_admin") or u.get("rol") in ("ingeniero", "analizador_materiales")) and (mismo or u.get("es_admin"))
    if not mismo:
        return False
    if etapa == "subgerente":
        return u.get("rol") == "subgerente"
    areas = _areas_responsabilidad(u)
    return u.get("rol") == "jefe" and (_norm(area) in areas or "TODAS" in areas)


def _visible(r, u):
    correo = u["correo"]
    return bool(u.get("es_admin") or correo in (r.solicitante_email, r.jefe_email, r.analista_email, r.subgerente_email))


def _dict(r):
    return {c.name: getattr(r, c.name) for c in r.__table__.columns} | {"datos": json.loads(r.datos_json), "historial": json.loads(r.historial_json)}


def listar(correo):
    with Session(engine) as s:
        u = _usuario(s, correo)
        q = select(SolicitudMaterial).order_by(SolicitudMaterial.id.desc())
        if not u.get("es_admin"):
            q = q.where(or_(SolicitudMaterial.solicitante_email == correo,
                           SolicitudMaterial.jefe_email == correo,
                           SolicitudMaterial.analista_email == correo,
                           SolicitudMaterial.subgerente_email == correo))
        return [_dict(r) for r in s.scalars(q).all()]


def _validar(d, tipo, completo=True):
    if tipo not in ("Stock de seguridad", "MRP"):
        raise ValueError("Tipo de solicitud inválido.")
    if completo:
        for k in ("material", "descripcion", "unidad", "centro", "area") + (("justificacion",) if tipo == "Stock de seguridad" else ()):
            if not str(d.get(k, "")).strip():
                raise ValueError(f"Debes completar {k}.")
        if d.get("criticidad") not in ("A", "B", "C", "Z", "M"):
            raise ValueError("Selecciona una criticidad válida.")
        if d.get("m_rotativo") and d.get("criticidad") != "M":
            raise ValueError("Un motor o bomba reparable rotativo debe conservar la categoría M.")
    if tipo == "MRP":
        if d.get("criticidad") == "M" or d.get("m_rotativo"):
            raise ValueError("Los materiales M no se pueden incorporar al MRP.")
        if completo:
            if d.get("perfil") == "Stock de seguridad" and float(d.get("cantidad", 0)) <= 0:
                raise ValueError("Indica la cantidad solicitada de stock de seguridad.")
            if d.get("perfil") not in PERFILES or d.get("mercado") not in ("Nacional", "Importado"):
                raise ValueError("Completa perfil y mercado.")
            lt = d.get("lead_time", 0)
            if not isinstance(lt, int) or isinstance(lt, bool) or lt < 1:
                raise ValueError("El lead time debe ser una cantidad entera de días mayor a cero.")
    if completo and tipo == "Stock de seguridad":
        if float(d.get("cantidad", 0)) <= 0:
            raise ValueError("La cantidad solicitada debe ser mayor a cero.")
    if completo and tipo == "Stock de seguridad" and d.get("solicitud_valorada"):
        if d.get("almacen") not in ("M010", "M100"):
            raise ValueError("Selecciona el almacén M010 o M100.")
        if float(d.get("precio_unitario", 0)) <= 0 or d.get("moneda") not in ("USD", "CLP"):
            raise ValueError("Indica precio unitario mayor a cero y moneda.")
        guia = d.get("guia") or {}
        calculada = criticidad(**{k: guia.get(k, False if k in ("m_rotativo", "estrategico") else None) for k in ("m_rotativo", "smac", "co", "te", "ma", "estrategico")})
        if d.get("criticidad") != calculada or bool(d.get("m_rotativo")) != bool(guia.get("m_rotativo")):
            raise ValueError("La criticidad debe coincidir con las respuestas de la guía.")
    if completo and d.get("uso_ss_registrado"):
        if not d.get("uso_ss_solo_emergencias"):
            raise ValueError("Confirma que el uso de SS incluye solo emergencias por falla y excluye trabajos planificados.")
        if float(d.get("periodo_consumo_meses",0)) <= 0:
            raise ValueError("Indica el período del registro de emergencias.")
        if float(d.get("uso_ss_cantidad",0)) > 0 and not str(d.get("uso_ss_respaldo", "")).strip():
            raise ValueError("Indica el aviso, OT correctiva o respaldo de las salidas por emergencia.")
    for k in ("precio_unitario", "cantidad", "ss_actual", "ss_propuesto", "precio", "stock", "reservas", "oc", "valor_bodega", "consumo_mensual", "stock_promedio", "desviacion", "lead_time_dias", "uso_ss_cantidad", "ss_historico_referencia", "periodo_consumo_meses"):
        if k in d and (not math.isfinite(float(d[k])) or float(d[k]) < 0):
            raise ValueError(f"Valor inválido para {k}.")


def _log(r, u, accion, comentario=""):
    h = json.loads(r.historial_json or "[]")
    h.append({"fecha": datetime.utcnow().isoformat(), "usuario": u["nombre"], "correo": u["correo"],
              "accion": accion, "estado": r.estado, "comentario": comentario,
              "version_datos": json.loads(r.datos_json or "{}")})
    r.historial_json = json.dumps(h, ensure_ascii=False)


def _aviso(s, correo, r, mensaje):
    s.add(NotificacionInterna(destinatario_email=correo, adf_id=None, tipo=f"materiales:{r.id}",
        titulo=f"Materiales #{r.id} · {r.estado}", mensaje=mensaje))


def guardar(correo, tipo, d, responsables, enviar=False, solicitud_id=None, version=None):
    _validar(d, tipo, completo=enviar)
    with Session(engine) as s:
        u = _usuario(s, correo)
        if not u.get("es_admin") and str(d.get("centro")) != str(u.get("centro")):
            raise ValueError("Solo puedes solicitar para tu centro.")
        if solicitud_id:
            r = s.get(SolicitudMaterial, solicitud_id, with_for_update=True)
            if not r or r.solicitante_email != correo or r.estado not in ("Borrador", "Devuelta"):
                raise ValueError("No puedes editar esta solicitud.")
            if r.version != version:
                raise ValueError("La solicitud cambió. Actualiza antes de continuar.")
            if r.tipo != tipo:
                raise ValueError("No puedes cambiar el tipo de una solicitud.")
        else:
            r = SolicitudMaterial(tipo=tipo, solicitante_email=correo, solicitante_nombre=u["nombre"],
                centro=str(d.get("centro", "")), area=d.get("area", ""), material=d.get("material", ""),
                estado="Borrador", historial_json="[]")
            s.add(r)
        if tipo == "Stock de seguridad":
            usuarios = [_a_dict(reg) for reg in s.scalars(select(UsuarioRootMine).where(UsuarioRootMine.activo.is_(True)))]
            responsables = responsables_stock(usuarios, d.get("centro", ""), d.get("area", ""))
            if enviar:
                if not responsables["jefe"]:
                    raise ValueError("No hay jefe responsable para este centro y área en el maestro ADF. Puedes guardar un borrador y revisar la asignación.")
                for etapa in ("analista", "subgerente"):
                    elegido = _usuario(s, responsables[etapa])
                    if not elegible(elegido, etapa, d["centro"], d["area"]):
                        raise ValueError(f"Revisa la cuenta fija de {etapa} en el maestro ADF.")
        if enviar:
            duplicada = s.scalar(select(SolicitudMaterial.id).where(
                SolicitudMaterial.centro == str(d["centro"]), SolicitudMaterial.material == d["material"],
                SolicitudMaterial.tipo == tipo,
                SolicitudMaterial.estado.in_(("Pendiente jefe", "Pendiente análisis", "Pendiente subgerente", "Pendiente carga SAP", PENDIENTE_MRP)),
                SolicitudMaterial.id != (r.id or -1)))
            if duplicada:
                raise ValueError(f"Ya existe la solicitud activa #{duplicada} para este material y centro.")
        r.centro = str(d.get("centro", "")); r.area = d.get("area", ""); r.material = d.get("material", "")
        # Cualquier reenvío invalida la evaluación y vuelve al jefe.
        r.datos_json = json.dumps(d, ensure_ascii=False)
        r.jefe_email = responsables.get("jefe", "") if tipo == "Stock de seguridad" else ""
        r.analista_email = responsables.get("analista", "") if tipo == "Stock de seguridad" else ""
        r.subgerente_email = responsables.get("subgerente", "") if tipo == "Stock de seguridad" else ""
        if tipo == "MRP":
            for campo in ("justificacion", "respaldo", "analisis"):
                d.pop(campo, None)
            r.datos_json = json.dumps(d, ensure_ascii=False)
        if enviar:
            r.estado = PENDIENTE_MRP if tipo == "MRP" else "Pendiente jefe"
            if not r.fecha_envio:
                r.fecha_envio = datetime.utcnow()
        s.flush()
        _log(r, u, "Enviada" if enviar else "Borrador guardado")
        if enviar and tipo == "Stock de seguridad":
            _aviso(s, r.jefe_email, r, f"{r.solicitante_nombre} solicita {r.tipo}: {r.material}. Revisa Gestión de Materiales.")
        try:
            s.commit()
        except StaleDataError:
            raise ValueError("La solicitud cambió. Actualiza e intenta nuevamente.") from None
        return r.id


def responsables_stock(usuarios, centro, area):
    jefe = _resolver_responsable(centro, area, "jefe", usuarios=usuarios)
    return {"jefe": jefe["correo"] if jefe else "", "analista": ANALISTA_FIJO, "subgerente": SUBGERENTE_FIJO}


def es_reemplazo_jefe(r, u):
    return (r["tipo"] == "Stock de seguridad" and r["estado"] == "Pendiente jefe"
            and u.get("correo") == ANALISTA_FIJO and bool(u.get("es_admin"))
            and str(u.get("centro", "")) == str(r["centro"])
            and u.get("correo") != r["jefe_email"])


def puede_actuar(r, u):
    if r["tipo"] == "MRP":
        return False
    if r["estado"] == "Pendiente carga SAP":
        return bool(u.get("es_admin"))
    if es_reemplazo_jefe(r, u):
        return True
    if r["estado"] == "Pendiente análisis" and (u.get("correo") != ANALISTA_FIJO or r["analista_email"] != ANALISTA_FIJO):
        return False
    if r["estado"] == "Pendiente subgerente" and (u.get("correo") != SUBGERENTE_FIJO or r["subgerente_email"] != SUBGERENTE_FIJO):
        return False
    campo = {"Pendiente jefe": "jefe_email", "Pendiente análisis": "analista_email",
             "Pendiente subgerente": "subgerente_email", "Pendiente carga SAP": "analista_email"}.get(r["estado"])
    if not campo:
        return False
    etapa = {"Pendiente jefe": "jefe", "Pendiente análisis": "analista",
             "Pendiente subgerente": "subgerente", "Pendiente carga SAP": "analista"}[r["estado"]]
    return u["correo"] == r[campo] and elegible(u, etapa, r["centro"], r["area"])


def decidir(correo, solicitud_id, version, accion, comentario="", analisis=None, carga=None):
    with Session(engine) as s:
        u = _usuario(s, correo)
        r = s.get(SolicitudMaterial, solicitud_id, with_for_update=True)
        if not r or not puede_actuar(_dict(r), u):
            raise ValueError("No eres el responsable de la etapa actual.")
        if r.version != version:
            raise ValueError("La solicitud cambió. Actualiza antes de continuar.")
        if es_reemplazo_jefe(_dict(r), u):
            if not comentario.strip():
                raise ValueError("Indica el motivo de ausencia del jefe para validar como reemplazo.")
            comentario = f"Validación como reemplazo del jefe {r.jefe_email} por ausencia. {comentario.strip()}"
        d = json.loads(r.datos_json)
        _validar(d, r.tipo)
        previo = r.estado
        if accion in ("Rechazar", "Devolver"):
            if not comentario.strip():
                raise ValueError("Debes indicar el motivo.")
            r.estado = "Rechazada" if accion == "Rechazar" else "Devuelta"
            destino = r.solicitante_email
        elif accion == "Aprobar" and previo == "Pendiente jefe":
            r.estado = "Pendiente análisis"; destino = r.analista_email
        elif accion == "Validar análisis" and previo == "Pendiente análisis":
            a = dict(analisis or {})
            if a.get("validacion_simple"):
                a.update(ss_propuesto=float(d.get("cantidad", 0)), criticidad=d["criticidad"], unidad=d.get("unidad", "unidades"))
            if r.tipo == "Stock de seguridad":
                if float(d.get("precio_unitario", 0)) <= 0:
                    raise ValueError("Devuelve la solicitud para que el solicitante informe el precio unitario. La validación solo agrega el stock valorado de la bodega.")
                a.update(precio=float(d["precio_unitario"]), moneda=d.get("moneda", "USD"),
                         ss_propuesto=float(d["cantidad"]), criticidad=d["criticidad"], unidad=d["unidad"])
            _validar({**d, **a}, r.tipo)
            if a.get("criticidad") == "M" and (r.tipo == "MRP" or d.get("perfil") in PERFILES):
                raise ValueError("Los materiales M no se incorporan al MRP.")
            if r.tipo == "Stock de seguridad" or d.get("perfil") == "Stock de seguridad":
                if float(a.get("ss_propuesto", 0)) <= 0 or float(a.get("precio", 0)) <= 0 or float(a.get("valor_bodega", 0)) <= 0:
                    raise ValueError("Completa stock propuesto, precio y valor total de bodega mayores a cero.")
                if not a.get("fecha_datos") or not a.get("moneda"):
                    raise ValueError("Indica fecha y moneda de la valorización.")
                a["impacto"] = impacto(a)
            d["analisis"] = a; r.estado = "Pendiente subgerente"; destino = r.subgerente_email
        elif accion == "Aprobar" and previo == "Pendiente subgerente":
            if not d.get("analisis"):
                raise ValueError("Falta la evaluación de materiales.")
            r.estado = "Pendiente carga SAP"; destino = r.analista_email
        elif accion in ("Cargado en SAP", "Marcar stock listo") and previo == "Pendiente carga SAP":
            c = dict(carga or {})
            if not c.get("referencia", "").strip() or not c.get("fecha"):
                raise ValueError("Indica la fecha real y referencia/evidencia de carga en SAP.")
            fecha = datetime.fromisoformat(c["fecha"])
            if fecha < r.fecha_envio or fecha > datetime.utcnow():
                raise ValueError("La fecha de carga debe estar entre el envío y el momento actual.")
            c["configuracion_aprobada"] = d.get("analisis", {})
            d["carga"] = c; r.fecha_carga = fecha; r.estado = LISTO_SS; destino = r.solicitante_email
        else:
            raise ValueError("Acción inválida para la etapa actual.")
        r.datos_json = json.dumps(d, ensure_ascii=False)
        _log(r, u, accion, comentario)
        mensaje = f"Solicitud #{r.id}, material {r.material}: {r.estado}. {comentario}"
        if r.fecha_carga:
            horas = (r.fecha_carga - r.fecha_envio).total_seconds() / 3600
            mensaje += f" Tiempo desde envío hasta SAP: {horas:.1f} horas."
        _aviso(s, destino, r, mensaje)
        if destino != r.solicitante_email:
            _aviso(s, r.solicitante_email, r, mensaje)
        try:
            s.commit()
        except StaleDataError:
            raise ValueError("Otra persona actualizó la solicitud. Actualiza la pantalla.") from None


def mrp_pendiente(r):
    """También recoge MRP enviados antes del cambio de flujo, sin borrar solicitudes."""
    return r["tipo"] == "MRP" and r["estado"] in (
        PENDIENTE_MRP, "Pendiente jefe", "Pendiente análisis", "Pendiente subgerente", "Pendiente carga SAP")


def _exportacion_datos(registros, lote, fecha, administrador):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Incorporacion MRP"
    columnas = ["Solicitud", "Material", "Descripción", "Unidad", "Centro", "Área", "Almacén", "Equipos",
                "Perfil MRP", "Mercado", "Lead time (días)", "Stock de seguridad solicitado", "Criticidad",
                "Fundamento criticidad", "Solicitante", "Correo solicitante", "Fecha registro", "Lote exportación"]
    ws.append(columnas)
    for r in registros:
        d = r["datos"]
        valores = [r["id"], d.get("material", r["material"]), d.get("descripcion", ""), d.get("unidad", ""),
                   r["centro"], r["area"], d.get("almacen", ""), d.get("equipos", ""), d.get("perfil", ""),
                   d.get("mercado", ""), d.get("lead_time", ""), d.get("cantidad", 0), d.get("criticidad", ""),
                   d.get("razon_criticidad", ""), r["solicitante_nombre"], r["solicitante_email"],
                   str(r["fecha_envio"] or r["fecha_creacion"]), lote]
        ws.append(valores)
        # Texto literal: conserva ceros iniciales y evita ejecutar fórmulas de entrada.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="163247")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx in range(1, len(columnas)+1):
        ws.column_dimensions[get_column_letter(idx)].width = 23 if idx not in (3,8,14) else 42
    meta = wb.create_sheet("Registro exportación")
    for row in [("Lote", lote), ("Fecha UTC", fecha), ("Administrador", administrador),
                ("Materiales", len(registros)),
                ("Estado", "Listo en RootMine por exportación; la carga efectiva en SAP se realiza externamente.")]:
        meta.append(row)
        for cell in meta[meta.max_row]:
            if isinstance(cell.value, str): cell.data_type = "s"
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 100
    contenido = BytesIO()
    wb.save(contenido)
    return contenido.getvalue()


def exportar_mrp(correo, seleccion, lote):
    """Genera el Excel al clic y cierra solo sus filas, en una transacción.

    seleccion es una lista de (id, versión) de la vista presentada. Un lote es
    idempotente: un reintento recupera su snapshot sin volver a cerrar filas.
    No usa session_state: la descarga diferida corre en otro hilo.
    """
    if not seleccion or len({n for n, _ in seleccion}) != len(seleccion):
        raise ValueError("No hay pendientes válidos para exportar.")
    from uuid import UUID
    UUID(lote)
    with Session(engine) as s:
        u = _usuario(s, correo)
        if not u.get("es_admin"):
            raise ValueError("Solo el administrador puede exportar los pendientes MRP.")
        registros = []
        for n, version in sorted(seleccion):
            r = s.get(SolicitudMaterial, n, with_for_update=True)
            if not r or r.tipo != "MRP":
                raise ValueError("El lote contiene una solicitud inválida.")
            registros.append((r, version))
        recuperados = []
        for r, _ in registros:
            exp = json.loads(r.datos_json).get("exportacion", {})
            if exp.get("lote") == lote:
                snapshot = next((h["version_datos"] for h in reversed(json.loads(r.historial_json))
                                 if h.get("accion") == "Exportado MRP" and h.get("version_datos", {}).get("exportacion", {}).get("lote") == lote), None)
                if snapshot is None:
                    raise ValueError("No se encontró el respaldo de este lote.")
                recuperados.append(_dict(r) | {"datos": snapshot})
        if recuperados:
            if len(recuperados) != len(registros):
                raise ValueError("El lote no coincide con el respaldo registrado.")
            exp = recuperados[0]["datos"]["exportacion"]
            return _exportacion_datos(recuperados, lote, exp["fecha"], exp["administrador"])
        fecha = datetime.utcnow()
        for r, version in registros:
            if r.version != version or not mrp_pendiente(_dict(r)):
                raise ValueError("Los pendientes cambiaron. Actualiza la vista antes de descargar.")
            _validar(json.loads(r.datos_json), "MRP")
        # Si falla la generación del Excel, no se marca ninguna solicitud.
        excel = _exportacion_datos([_dict(r) for r, _ in registros], lote, fecha.isoformat(), correo)
        for r, _ in registros:
            d = json.loads(r.datos_json)
            d["exportacion"] = {"lote": lote, "fecha": fecha.isoformat(), "administrador": correo}
            r.datos_json = json.dumps(d, ensure_ascii=False)
            r.estado = LISTO_MRP
            # fecha_carga se reserva para una carga SAP efectiva, nunca para un Excel.
            _log(r, u, "Exportado MRP", "Listo por exportación a Excel para incorporación al MRP.")
            _aviso(s, r.solicitante_email, r, f"Material {r.material}: exportado por administración en el lote {lote}.")
        try:
            s.commit()
        except StaleDataError:
            raise ValueError("Los pendientes cambiaron. Actualiza la vista antes de descargar.") from None
        return excel


def recuperar_lote_mrp(correo, lote):
    with Session(engine) as s:
        u = _usuario(s, correo)
        if not u.get("es_admin"):
            raise ValueError("Solo el administrador puede recuperar un lote MRP.")
        filas = []
        for r in s.scalars(select(SolicitudMaterial).where(SolicitudMaterial.tipo == "MRP").order_by(SolicitudMaterial.id)):
            for h in reversed(json.loads(r.historial_json)):
                d = h.get("version_datos", {})
                if h.get("accion") == "Exportado MRP" and d.get("exportacion", {}).get("lote") == lote:
                    filas.append(_dict(r) | {"datos": d})
                    break
        if not filas:
            raise ValueError("Lote no encontrado.")
        exp = filas[0]["datos"]["exportacion"]
        return _exportacion_datos(filas, lote, exp["fecha"], exp["administrador"])


def sincronizar_flujo_stock(correo):
    """Actualiza solicitudes abiertas al maestro ADF; conserva cierres y aprobaciones pasadas."""
    with Session(engine) as s:
        u = _usuario(s, correo)
        if not u.get("es_admin"):
            raise ValueError("Solo administración puede actualizar las asignaciones del flujo.")
        usuarios = [_a_dict(reg) for reg in s.scalars(select(UsuarioRootMine).where(UsuarioRootMine.activo.is_(True)))]
        filas = s.scalars(select(SolicitudMaterial).where(
            SolicitudMaterial.tipo == "Stock de seguridad",
            SolicitudMaterial.estado.in_(("Borrador", "Devuelta", "Pendiente jefe", "Pendiente análisis", "Pendiente subgerente", "Pendiente carga SAP"))
        ).order_by(SolicitudMaterial.id).with_for_update()).all()
        for r in filas:
            ruta = responsables_stock(usuarios, r.centro, r.area)
            cambios = []
            for etapa in ("jefe", "analista", "subgerente"):
                if etapa == "jefe" and (r.estado not in ("Borrador", "Devuelta", "Pendiente jefe") or not ruta[etapa]):
                    continue
                campo = etapa + "_email"
                anterior = getattr(r, campo)
                if anterior != ruta[etapa]:
                    setattr(r, campo, ruta[etapa])
                    cambios.append(f"{etapa}: {anterior or 'sin asignación'} → {ruta[etapa]}")
            if cambios:
                _log(r, u, "Responsables actualizados", "; ".join(cambios))
                destino = {"Pendiente jefe": r.jefe_email, "Pendiente análisis": r.analista_email,
                           "Pendiente subgerente": r.subgerente_email}.get(r.estado)
                if destino:
                    _aviso(s, destino, r, f"Solicitud #{r.id}: revisa la etapa {r.estado}. Responsables actualizados según maestro ADF.")
        try:
            s.commit()
        except StaleDataError:
            raise ValueError("Las solicitudes cambiaron. Actualiza la vista.") from None
