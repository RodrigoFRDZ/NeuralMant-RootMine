"""Mi Gestión: panel compacto, avances, evidencias y reportes sin duplicados."""
import io
from datetime import date, datetime
import pandas as pd
import streamlit as st
import plotly.graph_objects as go
from database import mi_gestion as db
from database import mi_gestion_adjuntos as storage
from database.repositorio_adf import listar_adf

ESTADOS = ['Pendiente','En ejecución','Finalizada','Postergada','Descartada']
CERRADOS = {'Finalizada','Descartada'}


def _dias(fecha):
    try: return (date.fromisoformat(str(fecha)[:10])-date.today()).days
    except (ValueError,TypeError): return None


def _avance_estado(estado, anterior=0):
    anterior = int(anterior or 0)
    if estado == 'Finalizada': return 100
    if estado in ('Pendiente','Descartada'): return 0
    if estado == 'En ejecución': return anterior if 1 <= anterior <= 99 else 10
    return min(99,max(0,anterior))


def _pct(r):
    return 100 if r['estado']=='Finalizada' else 0 if r['estado'] in ('Pendiente','Descartada') else int(r.get('avance') or 0)


def _excel(hojas):
    b=io.BytesIO()
    with pd.ExcelWriter(b,engine='openpyxl') as w:
        for nombre,df in hojas.items():
            df.to_excel(w,sheet_name=nombre[:31],index=False)
            ws=w.sheets[nombre[:31]]
            ws.freeze_panes='A2'
            ws.auto_filter.ref=ws.dimensions
            from openpyxl.styles import Font,PatternFill
            for c in ws[1]:
                c.fill=PatternFill('solid',fgColor='20354A')
                c.font=Font(color='FFFFFF',bold=True)
            for col in ws.columns:
                letter=col[0].column_letter
                ws.column_dimensions[letter].width=min(58,max(14,max(len(str(c.value or '')) for c in list(col)[:150])+2))
    return b.getvalue()


def _tabla(registros, objetivos):
    nombres={o['id']:o['titulo'] for o in objetivos}
    return pd.DataFrame([{
        'Objetivo':nombres.get(r.get('padre_id'),'Gestión independiente'),
        'Actividad':r['titulo'],'Estado':r['estado'],'Avance (%)':_pct(r),
        'Responsable':r.get('responsable') or '',
        'Fecha compromiso':r.get('fecha_limite') or '',
        'Días para vencer':_dias(r.get('fecha_limite')),
        'Última gestión':r.get('ultima_gestion') or '',
        'Presentada':'Sí' if r.get('presentada') else 'No',
    } for r in registros])


def _grafico(tareas, objetivos):
    import streamlit as st
    datos=[('Vencidas',sum((_dias(t.get('fecha_limite')) or 0)<0 for t in tareas),'#E26464'),
           ('Por vencer',sum(_dias(t.get('fecha_limite')) is not None and 0<=_dias(t.get('fecha_limite'))<=7 for t in tareas),'#E7AD57'),
           ('En plazo',sum(_dias(t.get('fecha_limite')) is not None and _dias(t.get('fecha_limite'))>7 for t in tareas),'#3EB79A'),
           ('Sin fecha',sum(_dias(t.get('fecha_limite')) is None for t in tareas),'#7F91A7')]
    fig=go.Figure(go.Bar(x=[x[1] for x in datos],y=[x[0] for x in datos],orientation='h',
                         marker_color=[x[2] for x in datos],text=[x[1] for x in datos],textposition='outside',
                         hovertemplate='%{y}: %{x} actividades<extra></extra>'))
    tema = str(st.get_option('theme.base') or '').lower()
    color_texto = '#596C80' if tema == 'light' else '#B7C4D3'
    fig.update_layout(height=240,margin=dict(l=5,r=35,t=10,b=10),paper_bgcolor='rgba(0,0,0,0)',
        plot_bgcolor='rgba(0,0,0,0)',font=dict(color=color_texto,size=12),showlegend=False,
        xaxis=dict(visible=False,range=[0,max([x[1] for x in datos]+[1])*1.2]),
        yaxis=dict(showgrid=False,autorange='reversed'),bargap=.45)
    st.plotly_chart(fig,use_container_width=True,config={'displayModeBar':False},key='grafico_estados_gestion')


def _resumen(tareas,objetivos,usuario):
    activas=[t for t in tareas if t['estado'] not in CERRADOS]
    vencidas=[t for t in activas if _dias(t.get('fecha_limite')) is not None and _dias(t.get('fecha_limite'))<0]
    proximas=[t for t in activas if _dias(t.get('fecha_limite')) is not None and 0<=_dias(t.get('fecha_limite'))<=7]
    estancadas=[]
    for t in activas:
        try:
            if (date.today()-date.fromisoformat(str(t.get('actualizado') or '')[:10])).days>=15:estancadas.append(t)
        except ValueError:pass
    avance=round(sum(_pct(t) for t in tareas if t['estado']!='Descartada')/max(1,sum(t['estado']!='Descartada' for t in tareas)))
    a,b,c,d,e=st.columns(5)
    a.metric('Avance de ejecución',f'{avance}%');b.metric('Activas',len(activas));c.metric('Vencidas',len(vencidas));d.metric('Por vencer (7 días)',len(proximas));e.metric('Sin actualizar (15 días)',len(estancadas))
    c1,c2=st.columns([1,1])
    with c1:
        st.markdown('**Situación de plazos**')
        _grafico(activas,objetivos)
    with c2:
        st.markdown('**Avance por objetivo**')
        avances=[]
        for o in objetivos:
            hijos=[t for t in tareas if t.get('padre_id')==o['id'] and t['estado']!='Descartada']
            if hijos:avances.append({'Objetivo':o['titulo'],'Avance (%)':round(sum(_pct(t) for t in hijos)/len(hijos))})
        if avances:
            df=pd.DataFrame(avances).sort_values('Avance (%)',ascending=False)
            st.dataframe(df,hide_index=True,use_container_width=True,height=240,column_config={'Avance (%)':st.column_config.ProgressColumn('Avance (%)',min_value=0,max_value=100,format='%d%%')})
        else:st.caption('Aún no hay tareas asociadas a objetivos.')
    st.caption('Los porcentajes reflejan ejecución de tareas, no cumplimiento de los indicadores operacionales.')
    filtro=st.segmented_control('Atención prioritaria',['Vencidas','Por vencer','Sin actualización','Todas'],default='Vencidas',key='filtro_prioridad')
    seleccion={'Vencidas':vencidas,'Por vencer':proximas,'Sin actualización':estancadas,'Todas':activas}[filtro]
    seleccion=sorted(seleccion,key=lambda t:(_dias(t.get('fecha_limite')) is None,_dias(t.get('fecha_limite')) if _dias(t.get('fecha_limite')) is not None else 99999))
    if seleccion:
        df=_tabla(seleccion,objetivos)
        st.dataframe(df[['Actividad','Objetivo','Responsable','Fecha compromiso','Días para vencer','Avance (%)']],
            hide_index=True,use_container_width=True,height=min(380,40+35*len(df)),
            column_config={'Avance (%)':st.column_config.ProgressColumn('Avance (%)',min_value=0,max_value=100,format='%d%%')})
        st.caption('Abre la tarea en «Objetivos y tareas» para actualizar su avance o seguimiento.')
        st.download_button('📥 Descargar prioridades',_excel({'Prioridades':df}),file_name='Prioridades_Mi_Gestion.xlsx',key='excel_prioridades')
    else:st.info('No hay actividades en esta categoría.')


def _evidencias(usuario,tarea):
    with st.expander('📎 Evidencias y respaldos',expanded=False):
        adjuntos=db.listar_adjuntos(usuario,tarea['id'])
        for a in adjuntos:
            cols=st.columns([5,1,1])
            cols[0].caption(('⭐ ' if a['principal'] else '')+a['nombre']+' · '+str(round(a['bytes']/1024))+' KB')
            if storage.configurado():
                if cols[1].button('Ver',key=f'ver_adj_{a["id"]}'):
                    try:
                        contenido=storage.descargar(a['ruta'])
                        st.download_button('Descargar '+a['nombre'],contenido,file_name=a['nombre'],key=f'bajar_{a["id"]}')
                        if a['tipo'].startswith('image/'):st.image(contenido,width=300)
                    except Exception as ex:st.error(str(ex))
            if cols[2].button('⭐',key=f'principal_{a["id"]}',help='Usar como evidencia principal'):
                db.marcar_principal(usuario,tarea['id'],a['id']);st.rerun()
        if not storage.configurado():
            st.caption('Para habilitar adjuntos, configura SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY y un bucket privado «mi-gestion-evidencias».')
            return
        if len(adjuntos)>=5:
            st.caption('Límite de cinco archivos por actividad.')
            return
        archivo=st.file_uploader('Adjuntar respaldo opcional (máx. 5 MB)',type=['jpg','jpeg','png','webp','pdf','xlsx','xls'],key=f'archivo_{tarea["id"]}')
        if archivo and st.button('Guardar respaldo',key=f'guardar_archivo_{tarea["id"]}'):
            datos=archivo.getvalue()
            if len(datos)>5*1024*1024:st.error('Archivo demasiado grande. Máximo 5 MB.');return
            tipo=archivo.type or 'application/octet-stream'
            try:
                if tipo.startswith('image/'):
                    from PIL import Image,ImageOps
                    img=ImageOps.exif_transpose(Image.open(io.BytesIO(datos))).convert('RGB')
                    img.thumbnail((1600,1600))
                    b=io.BytesIO();img.save(b,format='JPEG',quality=78,optimize=True)
                    datos=b.getvalue();tipo='image/jpeg';nombre=archivo.name.rsplit('.',1)[0]+'.jpg'
                else:nombre=archivo.name
                ruta=storage.subir(db.OWNER,tarea['id'],nombre,datos,tipo)
                try:db.guardar_adjunto(usuario,tarea['id'],nombre,ruta,tipo,len(datos))
                except Exception:
                    storage.eliminar(ruta);raise
                st.rerun()
            except Exception as ex:st.error('No se pudo guardar el respaldo: '+str(ex))


def _tarea(usuario,t,objetivo=None):
    cols=st.columns([5,2])
    cols[0].markdown(f'**{t["titulo"]}**')
    cols[0].caption(f'{t.get("finalidad") or "Sin finalidad"} · {t.get("responsable") or "Sin responsable"} · Plazo: {t.get("fecha_limite") or "Sin fecha"}')
    nuevo=cols[1].selectbox('Estado',ESTADOS,index=ESTADOS.index(t['estado']) if t['estado'] in ESTADOS else 0,key=f'estado_{t["id"]}')
    if nuevo!=t['estado']:
        db.actualizar(usuario,t['id'],estado=nuevo,avance=_avance_estado(nuevo,t.get('avance')))
        st.rerun()
    if nuevo=='En ejecución':
        p=cols[1].number_input('Avance (%)',min_value=1,max_value=99,value=max(1,min(99,int(t.get('avance') or 10))),key=f'pct_{t["id"]}')
        if p!=int(t.get('avance') or 0) and cols[1].button('Guardar avance',key=f'guardar_pct_{t["id"]}'):
            db.actualizar(usuario,t['id'],avance=p);st.rerun()
    with st.expander('Seguimiento y detalles',expanded=False):
        with st.form(f'nota_{t["id"]}'):
            nota=st.text_area('Última gestión realizada',value=t.get('ultima_gestion') or '')
            if st.form_submit_button('Guardar seguimiento'):
                db.actualizar(usuario,t['id'],ultima_gestion=nota);st.rerun()
        _evidencias(usuario,t)


def mostrar_mi_gestion():
    usuario=st.session_state.get('usuario_actual') or {}
    if not db.autorizado(usuario):
        st.error('No tienes acceso a este módulo.');st.stop()
    db.preparar()
    registros=db.listar(usuario)
    objetivos=[r for r in registros if r['tipo']=='objetivo']
    tareas=[r for r in registros if r['tipo'] in ('tarea','compromiso')]
    st.title('🔒 Mi Gestión')
    st.caption('Seguimiento privado de iniciativas, ejecución y resultados')
    tab0,tab1,tab2,tab3=st.tabs(['📊 Resumen ejecutivo','🎯 Objetivos y tareas','✨ Recomendaciones','📈 Resultados y reportes'])
    with tab0:_resumen(tareas,objetivos,usuario)
    with tab1:
        with st.expander('➕ Nuevo objetivo de mejora'):
            with st.form('nuevo_objetivo',clear_on_submit=True):
                titulo=st.text_input('Nombre del objetivo',key='obj_titulo')
                alcance=st.text_area('Justificación / alcance')
                indicador=st.text_input('Indicador')
                a,b,c=st.columns(3)
                base=a.text_input('Línea base');meta=b.text_input('Meta');actual=c.text_input('Valor actual')
                if st.form_submit_button('Guardar objetivo'):
                    if titulo.strip():
                        db.guardar(usuario,'objetivo',titulo.strip(),alcance,indicador=indicador,linea_base=base,meta=meta,valor_actual=actual)
                        st.success('Objetivo guardado.');st.rerun()
                    else:st.warning('Ingresa el nombre del objetivo.')
        modo=st.segmented_control('Mostrar tareas',['Activas','Finalizadas por presentar','Archivadas','Todas'],default='Activas')
        for o in objetivos:
            hijos=[t for t in tareas if t.get('padre_id')==o['id']]
            if modo=='Activas':visibles=[t for t in hijos if t['estado'] not in CERRADOS]
            elif modo=='Finalizadas por presentar':visibles=[t for t in hijos if t['estado']=='Finalizada' and not t.get('presentada')]
            elif modo=='Archivadas':visibles=[t for t in hijos if t['estado']=='Finalizada' and t.get('presentada')]
            else:visibles=hijos
            validas=[t for t in hijos if t['estado']!='Descartada']
            avance=round(sum(_pct(t) for t in validas)/len(validas)) if validas else 0
            with st.expander(f'🎯 {o["titulo"]} · {avance}% de ejecución',expanded=False):
                st.caption(f'{o.get("objetivo") or "Sin descripción"} · {o.get("indicador") or "Sin indicador"} · Base: {o.get("linea_base") or "—"} · Actual: {o.get("valor_actual") or "—"} · Meta: {o.get("meta") or "—"}')
                st.progress(avance/100,text=f'Avance de tareas: {avance}%')
                for t in visibles:
                    if t['estado']=='Finalizada' and not t.get('presentada'):
                        a,b=st.columns([5,2]);a.markdown(f'**{t["titulo"]}**');a.caption('Finalizada · pendiente de informar')
                        if b.button('✓ Ya fue presentado',key=f'presentada_{t["id"]}'):
                            db.actualizar(usuario,t['id'],presentada=1);st.rerun()
                        _evidencias(usuario,t)
                    elif t['estado']=='Finalizada' and t.get('presentada'):
                        st.caption('✓ Presentada · '+t['titulo'])
                        _evidencias(usuario,t)
                    else:_tarea(usuario,t,o)
                with st.form(f'crear_tarea_{o["id"]}',clear_on_submit=True):
                    st.markdown('**Agregar actividad**')
                    titulo=st.text_input('Actividad',key=f'nueva_tarea_{o["id"]}')
                    finalidad=st.text_input('Finalidad')
                    resp=st.text_input('Responsable')
                    plazo=st.date_input('Fecha compromiso',value=None)
                    if st.form_submit_button('Agregar actividad'):
                        if titulo.strip():
                            db.guardar(usuario,'tarea',titulo.strip(),finalidad=finalidad,padre_id=o['id'],responsable=resp,fecha_limite=plazo.isoformat() if plazo else '')
                            st.success('Actividad guardada.');st.rerun()
                        else:st.warning('Ingresa una actividad.')
                with st.form(f'actualizar_obj_{o["id"]}'):
                    actual=st.text_input('Valor actual del indicador',value=o.get('valor_actual') or '')
                    estado=st.selectbox('Estado del objetivo',ESTADOS,index=ESTADOS.index(o['estado']) if o['estado'] in ESTADOS else 0)
                    if st.form_submit_button('Actualizar objetivo'):
                        db.actualizar(usuario,o['id'],valor_actual=actual,estado=estado);st.rerun()
        independientes=[t for t in tareas if t.get('padre_id') is None]
        with st.expander('➕ Actividad independiente (sin objetivo)'):
            with st.form('nueva_independiente',clear_on_submit=True):
                titulo=st.text_input('Actividad o gestión puntual')
                finalidad=st.text_input('Finalidad')
                resp=st.text_input('Responsable')
                plazo=st.date_input('Fecha compromiso',value=None)
                if st.form_submit_button('Agregar') and titulo.strip():
                    db.guardar(usuario,'tarea',titulo.strip(),finalidad=finalidad,responsable=resp,fecha_limite=plazo.isoformat() if plazo else '');st.rerun()
        if independientes:
            with st.expander('Actividades independientes'):
                for t in independientes:
                    if modo=='Activas' and t['estado'] in CERRADOS:continue
                    if modo=='Finalizadas por presentar' and (t['estado']!='Finalizada' or t.get('presentada')):continue
                    if modo=='Archivadas' and (t['estado']!='Finalizada' or not t.get('presentada')):continue
                    if t['estado']=='Finalizada' and not t.get('presentada'):
                        st.write(t['titulo'])
                        if st.button('✓ Ya fue presentado',key=f'present_ind_{t["id"]}'):
                            db.actualizar(usuario,t['id'],presentada=1);st.rerun()
                        _evidencias(usuario,t)
                    elif t['estado']=='Finalizada':st.caption('✓ '+t['titulo'])
                    else:_tarea(usuario,t)
    with tab2:
        st.info('Las sugerencias requieren aprobación. No se crean tareas automáticamente.')
        existentes={r['origen'] for r in registros if r['tipo']=='recomendacion'}
        propuestas=[]
        for a in listar_adf():
            if str(getattr(a,'estado','')).lower()=='borrador':continue
            plan=getattr(a,'plan_prevencion','') or ''
            if plan and plan not in ('[]','{}'):
                propuestas.append((f'ADF #{a.id}',f'Revisar efectividad de acciones del ADF #{a.id} · {a.equipo}','Comprobar ejecución y resultados'))
            if len(propuestas)>=15:break
        for origen,titulo,finalidad in propuestas:
            if origen in existentes:continue
            with st.container(border=True):
                st.write('**'+titulo+'**');st.caption(finalidad)
                if st.button('Guardar para evaluar',key='rec_'+origen):
                    db.guardar(usuario,'recomendacion',titulo,finalidad=finalidad,origen=origen);st.rerun()
        for r in registros:
            if r['tipo']!='recomendacion' or r['estado'] in CERRADOS:continue
            with st.container(border=True):
                st.write(r['titulo']);st.caption(f'Origen: {r["origen"]}')
                a,b=st.columns(2)
                if a.button('Convertir en actividad',key=f'aceptar_{r["id"]}'):
                    db.guardar(usuario,'tarea',r['titulo'],finalidad=r['finalidad'],origen=r['origen'])
                    db.actualizar(usuario,r['id'],estado='Finalizada');st.rerun()
                if b.button('Descartar',key=f'descartar_{r["id"]}'):
                    db.actualizar(usuario,r['id'],estado='Descartada');st.rerun()
    with tab3:
        st.markdown('**Gestión para el próximo informe**')
        incluir=[t for t in tareas if t['estado']!='Descartada' and not (t['estado']=='Finalizada' and t.get('presentada'))]
        if incluir:
            df=_tabla(incluir,objetivos)
            st.dataframe(df,hide_index=True,use_container_width=True)
            st.download_button('📥 Descargar Mi Gestión actual',_excel({'Actividades':df}),file_name='Mi_Gestion_RootMine.xlsx',key='descargar_gestion')
        else:st.info('No hay actividades nuevas o pendientes de presentar.')
        pendientes=[t for t in tareas if t['estado']=='Finalizada' and not t.get('presentada')]
        st.markdown(f'**Finalizadas por presentar: {len(pendientes)}**')
        for t in pendientes:
            a,b=st.columns([5,2]);a.caption(t['titulo'])
            if b.button('✓ Ya fue presentado',key=f'present_report_{t["id"]}'):
                db.actualizar(usuario,t['id'],presentada=1);st.rerun()
        st.caption('Una tarea presentada queda archivada y no aparece en futuras descargas. Los registros y evidencias se conservan.')
