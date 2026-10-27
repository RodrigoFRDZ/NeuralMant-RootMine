import sys, tempfile, unittest
from io import BytesIO
from uuid import uuid4
from unittest.mock import patch
from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime, timedelta
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from database.modelos import Base, UsuarioRootMine, NotificacionInterna
from database.modelos_materiales import SolicitudMaterial
import database.materiales as m

class MaterialesTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.prev=m.engine
        m.engine=create_engine('sqlite:///'+str(Path(self.tmp.name)/'prueba.db'))
        Base.metadata.create_all(m.engine)
        with Session(m.engine) as s:
            for correo,rol,area in [('sol@test.cl','tecnico','FAENA'),('jefe@test.cl','jefe','FAENA'),('rfernandezc@agrosuper.com','analizador_materiales','TODAS'),('cvelasquez@agrosuper.com','subgerente','TODAS'),('otro@test.cl','tecnico','FAENA')]:
                s.add(UsuarioRootMine(nombre=correo,correo=correo,rol=rol,centro='1802',area=area,activo=True,es_admin=(correo=='rfernandezc@agrosuper.com'),responsable_de='[]'))
            s.commit()
        self.d=dict(precio_unitario=100,moneda='USD',material='123',descripcion='Rodamiento',unidad='UN',centro='1802',area='FAENA',justificacion='Detiene la línea si falta',criticidad='A',cantidad=2,m_rotativo=False,perfil='Stock de seguridad',mercado='Nacional',lead_time=10)
        self.resp={'jefe':'jefe@test.cl','analista':'rfernandezc@agrosuper.com','subgerente':'cvelasquez@agrosuper.com'}
    def tearDown(self):
        m.engine.dispose();m.engine=self.prev;self.tmp.cleanup()
    def row(self,n,correo='sol@test.cl'):
        return next(r for r in m.listar(correo) if r['id']==n)
    def action(self,n,correo,accion,**kw):
        m.decidir(correo,n,self.row(n,correo)['version'],accion,**kw)
    def test_solicitud_anterior_sin_precio_no_se_completa_en_administracion(self):
        d=dict(self.d)
        d.pop('precio_unitario')
        n=m.guardar('sol@test.cl','Stock de seguridad',d,self.resp,enviar=True)
        self.action(n,'jefe@test.cl','Aprobar')
        with self.assertRaisesRegex(ValueError,'solicitante'):
            self.action(n,m.ANALISTA_FIJO,'Validar análisis',analisis=dict(validacion_simple=True,precio=100,valor_bodega=1000,fecha_datos='2026-10-07',moneda='USD'))
        self.assertEqual(self.row(n)['estado'],'Pendiente análisis')
        self.action(n,m.ANALISTA_FIJO,'Devolver',comentario='Completar precio unitario')
        self.assertEqual(self.row(n)['estado'],'Devuelta')

    def test_precio_solicitante_y_guia_se_conservan_en_el_flujo(self):
        d=self.d | dict(solicitud_valorada=True,precio_unitario=25,moneda='CLP',almacen='M010',
            guia=dict(m_rotativo=False,estrategico=False,smac=3,co=1,te=2,ma=2))
        for cambio in [dict(precio_unitario=0),dict(almacen='XYZ'),dict(criticidad='C'),dict(guia={})]:
            with self.assertRaises(ValueError):m.guardar('sol@test.cl','Stock de seguridad',d|cambio,self.resp,enviar=True)
        n=m.guardar('sol@test.cl','Stock de seguridad',d,self.resp,enviar=True)
        self.assertEqual(self.row(n,'jefe@test.cl')['datos']['precio_unitario'],25)
        self.action(n,'jefe@test.cl','Aprobar')
        self.action(n,m.ANALISTA_FIJO,'Validar análisis',analisis=dict(validacion_simple=True,precio=999,moneda='USD',valor_bodega=1000,fecha_datos='2026-10-07'))
        a=self.row(n,m.SUBGERENTE_FIJO)['datos']['analisis']
        self.assertEqual(a['precio'],25)
        self.assertEqual(a['moneda'],'CLP')
        self.assertEqual(a['impacto']['valor_ss_total'],50)
        self.assertEqual(a['impacto']['porcentaje_ss_total_bodega_actual'],5)
        self.action(n,m.SUBGERENTE_FIJO,'Aprobar')
        self.assertEqual(self.row(n)['datos']['analisis'],a)

    def test_validacion_sin_comentario_conserva_datos_para_subgerente(self):
        n=m.guardar('sol@test.cl','Stock de seguridad',self.d,self.resp,enviar=True)
        self.action(n,'jefe@test.cl','Aprobar')
        a=dict(validacion_simple=True,ss_propuesto=99,criticidad='C',precio=100,valor_bodega=10000,fecha_datos='2026-10-07',moneda='USD')
        self.action(n,m.ANALISTA_FIJO,'Validar análisis',analisis=a)
        r=self.row(n,m.SUBGERENTE_FIJO)
        guardado=r['datos']['analisis']
        self.assertEqual(guardado['ss_propuesto'],2)
        self.assertEqual(guardado['criticidad'],'A')
        self.assertEqual(guardado['impacto']['valor_ss_total'],200)
        self.assertEqual(guardado['impacto']['bodega_con_material_solicitado'],10200)
        self.assertEqual(r['historial'][-1]['comentario'],'')
        self.action(n,m.SUBGERENTE_FIJO,'Aprobar')
        self.assertEqual(self.row(n)['datos']['analisis'],guardado)
        self.assertEqual(self.row(n)['estado'],'Pendiente carga SAP')

    def test_flujo_y_notificaciones(self):
        n=m.guardar('sol@test.cl','Stock de seguridad',self.d,self.resp,enviar=True)
        self.assertEqual(m.listar('otro@test.cl'),[])
        with self.assertRaises(ValueError): self.action(n,'sol@test.cl','Aprobar')
        self.action(n,'jefe@test.cl','Aprobar')
        a=dict(criticidad='A',ss_propuesto=2,ss_actual=1,stock=0,reservas=0,oc=1,precio=100,valor_bodega=10000,fecha_datos='2026-10-02',moneda='USD',fundamento='Cobertura mínima para continuidad')
        self.action(n,'rfernandezc@agrosuper.com','Validar análisis',analisis=a)
        self.assertEqual(self.row(n)['datos']['analisis']['impacto']['incremento_ss'],100)
        self.action(n,'cvelasquez@agrosuper.com','Aprobar')
        with self.assertRaises(ValueError): self.action(n,'rfernandezc@agrosuper.com','Cargado en SAP',carga={'fecha':'2099-01-01','referencia':'SAP'})
        self.action(n,'rfernandezc@agrosuper.com','Cargado en SAP',carga={'fecha':datetime.utcnow().isoformat(),'referencia':'SAP MM02 evidencia 123'})
        r=self.row(n);self.assertEqual(r['estado'],m.LISTO_SS);self.assertIsNotNone(r['fecha_carga']);self.assertEqual(len(r['historial']),5)
        with Session(m.engine) as s:
            avisos=s.scalars(select(NotificacionInterna)).all()
            self.assertTrue(any(a.destinatario_email=='sol@test.cl' and 'SAP' in a.mensaje for a in avisos))
        with self.assertRaises(ValueError): self.action(n,'rfernandezc@agrosuper.com','Cargado en SAP')
    def test_mrp_y_categoria_m(self):
        for campo in ['criticidad','m_rotativo']:
            d=self.d|{campo:'M' if campo=='criticidad' else True}
            for enviar in [False,True]:
                with self.assertRaises(ValueError):m.guardar('sol@test.cl','MRP',d,self.resp,enviar=enviar)
        for lt in [0,-1,1.5]:
            with self.assertRaises(ValueError):m.guardar('sol@test.cl','MRP',self.d|{'lead_time':lt},self.resp,enviar=True)
        n=m.guardar('sol@test.cl','MRP',self.d|{'perfil':'Pronóstico'},self.resp,enviar=True)
        self.assertEqual(self.row(n)['estado'],m.PENDIENTE_MRP)
        self.assertFalse(m.puede_actuar(self.row(n), {'correo':'rfernandezc@agrosuper.com','es_admin':True}))
    def test_devolver_reenviar_version_y_duplicados(self):
        n=m.guardar('sol@test.cl','Stock de seguridad',self.d,self.resp,enviar=True)
        with self.assertRaises(ValueError):m.guardar('sol@test.cl','Stock de seguridad',self.d,self.resp,enviar=True)
        r=self.row(n)
        with self.assertRaises(ValueError):m.decidir('jefe@test.cl',n,r['version']-1,'Aprobar')
        with self.assertRaises(ValueError):self.action(n,'jefe@test.cl','Rechazar')
        self.action(n,'jefe@test.cl','Devolver',comentario='Falta evidencia')
        r=self.row(n)
        m.guardar('sol@test.cl','Stock de seguridad',self.d|{'respaldo':'OT123'},self.resp,enviar=True,solicitud_id=n,version=r['version'])
        self.assertEqual(self.row(n)['estado'],'Pendiente jefe')
        self.assertEqual(self.row(n)['fecha_envio'],r['fecha_envio'])
    def test_arbol_e_impacto(self):
        self.assertEqual(m.criticidad(True,3,3,3,3,False),'M')
        self.assertEqual(m.criticidad(False,1,3,3,3,True),'Z')
        self.assertEqual(m.criticidad(False,3,1,3,3,False),'A')
        self.assertEqual(m.criticidad(False,3,3,1,1,False),'A')
        self.assertEqual(m.criticidad(False,3,3,2,2,False),'B')
        self.assertEqual(m.criticidad(False,3,3,2,3,False),'C')
        self.assertEqual(m.criticidad(False,3,3,3,1,False),'C')
        self.assertIsNone(m.impacto({})['porcentaje'])
        with self.assertRaises(ValueError):m.impacto({'precio':float('nan')})
    def test_responsables_y_centro(self):
        n=m.guardar('sol@test.cl','Stock de seguridad',self.d,self.resp|{'jefe':'otro@test.cl'},enviar=True)
        self.assertEqual(self.row(n)['jefe_email'],'jefe@test.cl')
        with self.assertRaises(ValueError):m.guardar('sol@test.cl','MRP',self.d|{'centro':'1901'},self.resp,enviar=True)


    def test_mrp_exportacion_incremental_y_recuperacion(self):
        d = {k:v for k,v in self.d.items() if k != 'justificacion'} | {'material':'000123', 'descripcion':'=NO_ES_FORMULA'}
        n=m.guardar('sol@test.cl','MRP',d,{},enviar=True)
        r=self.row(n)
        self.assertEqual(r['jefe_email'],'')
        self.assertNotIn('justificacion',r['datos'])
        lote=str(uuid4())
        with self.assertRaises(ValueError):m.exportar_mrp('sol@test.cl',[(n,r['version'])],lote)
        contenido=m.exportar_mrp('rfernandezc@agrosuper.com',[(n,r['version'])],lote)
        ws=load_workbook(BytesIO(contenido)).worksheets[0]
        self.assertEqual(ws.max_row,2)
        self.assertEqual(ws['B2'].value,'000123')
        self.assertEqual(ws['C2'].data_type,'s')
        listo=self.row(n)
        self.assertEqual(listo['estado'],m.LISTO_MRP)
        self.assertIsNone(listo['fecha_carga'])
        self.assertFalse(m.mrp_pendiente(listo))
        n2=m.guardar('sol@test.cl','MRP',self.d|{'material':'456'},{},enviar=True)
        reintento=m.exportar_mrp('rfernandezc@agrosuper.com',[(n,r['version'])],lote)
        self.assertEqual(load_workbook(BytesIO(reintento)).worksheets[0].max_row,2)
        self.assertTrue(m.mrp_pendiente(self.row(n2)))
        recuperado=m.recuperar_lote_mrp('rfernandezc@agrosuper.com',lote)
        self.assertEqual(load_workbook(BytesIO(recuperado)).worksheets[0]['B2'].value,'000123')
        self.assertEqual(len(self.row(n)['historial']),2)
        n2r=self.row(n2)
        segunda=m.exportar_mrp('rfernandezc@agrosuper.com',[(n2,n2r['version'])],str(uuid4()))
        self.assertEqual(load_workbook(BytesIO(segunda)).worksheets[0]['B2'].value,'456')
        with self.assertRaises(ValueError):m.recuperar_lote_mrp('otro@test.cl',lote)

    def test_exportacion_atomica_y_snapshot_obsoleto(self):
        n=m.guardar('sol@test.cl','MRP',self.d,{},enviar=True)
        n2=m.guardar('sol@test.cl','MRP',self.d|{'material':'456'},{},enviar=True)
        seleccion=[(n,self.row(n)['version']), (n2,self.row(n2)['version'])]
        with patch.object(m,'_exportacion_datos',side_effect=RuntimeError('Error archivo')):
            with self.assertRaises(RuntimeError):m.exportar_mrp('rfernandezc@agrosuper.com',seleccion,str(uuid4()))
        self.assertTrue(all(m.mrp_pendiente(self.row(i)) for i in [n,n2]))
        with self.assertRaises(ValueError):m.exportar_mrp('rfernandezc@agrosuper.com',[(n,999),seleccion[1]],str(uuid4()))
        self.assertTrue(all(m.mrp_pendiente(self.row(i)) for i in [n,n2]))
        m.exportar_mrp('rfernandezc@agrosuper.com',[seleccion[0]],str(uuid4()))
        with self.assertRaises(ValueError):m.exportar_mrp('rfernandezc@agrosuper.com',seleccion,str(uuid4()))
        self.assertTrue(m.mrp_pendiente(self.row(n2)))

    def test_mrp_legacy_y_duplicados(self):
        n=m.guardar('sol@test.cl','MRP',self.d,{},enviar=True)
        with self.assertRaises(ValueError):m.guardar('sol@test.cl','MRP',self.d,{},enviar=True)
        with Session(m.engine) as s:
            r=s.get(SolicitudMaterial,n);r.estado='Pendiente jefe';s.commit()
        r=self.row(n)
        self.assertTrue(m.mrp_pendiente(r))
        m.exportar_mrp('rfernandezc@agrosuper.com',[(n,r['version'])],str(uuid4()))
        self.assertEqual(self.row(n)['estado'],m.LISTO_MRP)

    def test_stock_cierre_solo_admin(self):
        n=m.guardar('sol@test.cl','Stock de seguridad',self.d,self.resp,enviar=True)
        self.action(n,'jefe@test.cl','Aprobar')
        self.action(n,'rfernandezc@agrosuper.com','Validar análisis',analisis=dict(criticidad='A',ss_propuesto=2,precio=100,valor_bodega=10000,fecha_datos='2026-10-02',moneda='USD',fundamento='Cobertura'))
        self.action(n,'cvelasquez@agrosuper.com','Aprobar')
        with self.assertRaises(ValueError):m.decidir('otro@test.cl',n,self.row(n)['version'],'Marcar stock listo',carga={'fecha':datetime.utcnow().isoformat(),'referencia':'SAP'})
        self.action(n,'rfernandezc@agrosuper.com','Marcar stock listo',carga={'fecha':datetime.utcnow().isoformat(),'referencia':'SAP MM02'})
        self.assertEqual(self.row(n)['estado'],m.LISTO_SS)

    def test_reemplazo_jefe_y_subgerente_obligatorio(self):
        n=m.guardar('sol@test.cl','Stock de seguridad',self.d,{},enviar=True)
        with self.assertRaises(ValueError):self.action(n,m.ANALISTA_FIJO,'Aprobar')
        self.action(n,m.ANALISTA_FIJO,'Aprobar',comentario='Jefe de vacaciones')
        r=self.row(n)
        self.assertEqual(r['estado'],'Pendiente análisis')
        self.assertIn('reemplazo del jefe',r['historial'][-1]['comentario'])
        self.assertIn('vacaciones',r['historial'][-1]['comentario'])
        self.action(n,m.ANALISTA_FIJO,'Validar análisis',analisis=dict(criticidad='A',ss_propuesto=2,precio=100,valor_bodega=10000,fecha_datos='2026-10-02',moneda='USD',fundamento='Cobertura'))
        for accion in ['Aprobar','Marcar stock listo']:
            with self.assertRaises(ValueError):self.action(n,m.ANALISTA_FIJO,accion,comentario='Subgerente ausente')
        self.assertEqual(self.row(n)['estado'],'Pendiente subgerente')
        self.action(n,m.SUBGERENTE_FIJO,'Aprobar')
        self.assertEqual(self.row(n)['estado'],'Pendiente carga SAP')

    def test_match_adf_multiarea_y_responsables_fijos(self):
        with Session(m.engine) as s:
            jefe=s.scalar(select(UsuarioRootMine).where(UsuarioRootMine.correo=='jefe@test.cl'))
            jefe.area='Jefatura';jefe.responsable_de='["FAENA", "ELABORADOS"]';s.commit()
        n=m.guardar('sol@test.cl','Stock de seguridad',self.d|{'area':'Elaborádos'}, {'jefe':'otro@test.cl','analista':'otro@test.cl','subgerente':'otro@test.cl'},enviar=True)
        r=self.row(n)
        self.assertEqual((r['jefe_email'],r['analista_email'],r['subgerente_email']),('jefe@test.cl',m.ANALISTA_FIJO,m.SUBGERENTE_FIJO))
        with self.assertRaises(ValueError):m.guardar('sol@test.cl','Stock de seguridad',self.d|{'area':'SIN JEFE','material':'555'},{},enviar=True)

    def test_sincronizacion_legacy_no_saltea_subgerente(self):
        n=m.guardar('sol@test.cl','Stock de seguridad',self.d,{},enviar=True)
        with Session(m.engine) as s:
            r=s.get(SolicitudMaterial,n);r.analista_email='otro@test.cl';r.subgerente_email='otro@test.cl';s.commit()
        m.sincronizar_flujo_stock(m.ANALISTA_FIJO)
        r=self.row(n)
        self.assertEqual(r['analista_email'],m.ANALISTA_FIJO)
        self.assertEqual(r['subgerente_email'],m.SUBGERENTE_FIJO)
        self.assertEqual(r['estado'],'Pendiente jefe')
        historial=len(r['historial'])
        m.sincronizar_flujo_stock(m.ANALISTA_FIJO)
        self.assertEqual(len(self.row(n)['historial']),historial)

    def test_valor_total_ss_cobertura_y_promedio(self):
        d=dict(ss_propuesto=6,ss_actual=2,precio=100,valor_bodega=10000,stock=4,reservas=0,oc=0,
               consumo_mensual=2,stock_promedio_observado=True,stock_promedio=8,uso_ss_registrado=True,uso_ss_cantidad=6,periodo_consumo_meses=3)
        i=m.impacto(d)
        self.assertEqual(i['valor_ss_total'],600)
        self.assertEqual(i['incremento_ss'],400)
        self.assertEqual(i['porcentaje_ss_total_bodega_actual'],6)
        self.assertEqual(i['valor_stock_material_actual'],400)
        self.assertEqual(i['reposicion'],200)
        self.assertEqual(i['bodega_proyectada'],10200)
        self.assertAlmostEqual(i['porcentaje_ss_total_bodega_proyectada'],600/10200*100)
        self.assertEqual(i['cobertura_ss_meses'],3)
        self.assertEqual(i['cobertura_ss_dias'],90)
        self.assertEqual(i['cobertura_stock_meses'],2)
        self.assertEqual(i['permanencia_promedio_meses'],4)
        sin=m.impacto(d|{'consumo_mensual':0,'valor_bodega':0,'stock_promedio_observado':False,'uso_ss_cantidad':0})
        self.assertIsNone(sin['cobertura_ss_meses'])
        self.assertIsNone(sin['permanencia_promedio_meses'])
        self.assertIsNone(sin['porcentaje_ss_total_bodega_actual'])
        self.assertIsNone(m.impacto({})['uso_ss_cantidad'])

    def test_uso_reserva_no_se_confunde_con_consumo_total(self):
        d=dict(precio=10,consumo_mensual=5,periodo_consumo_meses=12,uso_ss_registrado=True,uso_ss_cantidad=15,ss_historico_referencia=5)
        i=m.impacto(d)
        self.assertEqual(i['uso_ss_valorizado'],150)
        self.assertEqual(i['uso_ss_porcentaje_consumo_periodo'],25)
        self.assertEqual(i['uso_ss_reservas_equivalentes'],3)
        self.assertIsNone(m.impacto(d|{'uso_ss_registrado':False})['uso_ss_valorizado'])
        self.assertIsNone(m.impacto(d|{'consumo_mensual':0})['uso_ss_porcentaje_consumo_periodo'])
        with self.assertRaises(ValueError):m.impacto(d|{'uso_ss_cantidad':-1})
        with self.assertRaises(ValueError):m.impacto(d|{'consumo_mensual':float('nan')})

    def test_cobertura_ss_usa_solo_emergencias_y_exige_respaldo(self):
        d=dict(ss_propuesto=6,consumo_mensual=100,periodo_consumo_meses=12,uso_ss_registrado=True,uso_ss_cantidad=12)
        self.assertEqual(m.impacto(d)['cobertura_ss_meses'],6)
        self.assertEqual(m.impacto(d|{'consumo_mensual':200})['cobertura_ss_meses'],6)
        self.assertIsNone(m.impacto(d|{'uso_ss_registrado':False})['cobertura_ss_meses'])
        a=self.d|d|{'uso_ss_solo_emergencias':False}
        with self.assertRaises(ValueError):m._validar(a,'Stock de seguridad')
        with self.assertRaises(ValueError):m._validar(a|{'uso_ss_solo_emergencias':True},'Stock de seguridad')
        m._validar(a|{'uso_ss_solo_emergencias':True,'uso_ss_respaldo':'Aviso SAP / OT correctiva'},'Stock de seguridad')

if __name__=='__main__':unittest.main()
